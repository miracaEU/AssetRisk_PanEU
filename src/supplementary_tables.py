"""
supplementary_tables.py

Build Supplementary Tables S1-S4 of the ERL paper from the code and the
exposure input files.

  S1  OSM tags per infrastructure system (from the object types and tag
      columns present in the exposure inputs; the OSM extraction itself is
      not part of this repository)
  S2  Number of assets (OSM features, split at LAU boundaries) per country
      and system
  S3  Vulnerability / fragility curve IDs actually applied per asset type,
      plus a legend of every ID
  S4  Maximum reconstruction costs (min / median / max) with units

Inputs:
  {exposure_dir}/{System}/{System}_{ISO2}.parquet    exposure inputs
  {exposure_dir}/_stats/{ISO2}/report_*.txt         extraction reports (geometry types)
  data/EQ_fragility.xlsx, data/Table_D2_Hazard_Fragility_and_Vulnerability_Curves_V1.1.0_conversions.xlsx

Usage (from repo root):
  uv run python src/supplementary_tables.py
  uv run python src/supplementary_tables.py --exposure-dir /path/to/Exposure_files --output-dir tables
  uv run --with python-docx python src/supplementary_tables.py      # also writes a .docx
"""

import argparse
import re
from pathlib import Path

import pandas as pd
import pyarrow.parquet as pq

import constants as C
import hazard_earthquake as HEQ
from run_pipeline import Config

REPO_ROOT = Path(__file__).parent.parent

FOLDERS = {"Roadway": "roads", "Railway": "rail", "Airports": "airports", "Ports": "ports",
           "Power": "power", "Telecom": "telecom", "Education": "education", "Healthcare": "healthcare"}
SYS = ["roads", "rail", "airports", "ports", "power", "telecom", "education", "healthcare"]
SYS_LABEL = {"roads": "Roads", "rail": "Railways", "airports": "Airports", "ports": "Ports", "power": "Power",
             "telecom": "Telecommunications", "education": "Education", "healthcare": "Healthcare"}
CODE_KEY = {"airports": "air"}   # key used in constants.py dictionaries
ISO2_NAME = {"AD": "Andorra", "AL": "Albania", "AT": "Austria", "BE": "Belgium", "BG": "Bulgaria", "CH": "Switzerland",
             "CY": "Cyprus", "CZ": "Czechia", "DE": "Germany", "DK": "Denmark", "EE": "Estonia", "EL": "Greece",
             "ES": "Spain", "FI": "Finland", "FR": "France", "GB": "United Kingdom", "HR": "Croatia", "HU": "Hungary",
             "IE": "Ireland", "IS": "Iceland", "IT": "Italy", "LI": "Liechtenstein", "LT": "Lithuania",
             "LU": "Luxembourg", "LV": "Latvia", "MK": "North Macedonia", "MT": "Malta", "NL": "Netherlands",
             "NO": "Norway", "PL": "Poland", "PT": "Portugal", "RO": "Romania", "RS": "Serbia", "SE": "Sweden",
             "SI": "Slovenia", "SK": "Slovakia"}

# OSM key=value used for selection (inferred from the tag columns in the inputs)
KEYS = {"roads": "highway", "rail": "railway", "airports": "aeroway", "power": "power",
        "education": "amenity", "healthcare": "amenity"}
PORT_TAGS = {"harbour": "seamark:type=harbour; landuse=harbour", "port": "industrial=port (landuse=industrial)"}
TEL_TAGS = {"mast": "man_made=mast", "tower": "man_made=tower + tower:type=communication|telecommunication",
            "communications_tower": "man_made=communications_tower"}
GEOM_UNIT = {"LineString": "€/m", "Polygon": "€/m²", "Point": "€/unit"}


# ---------------------------------------------------------------------------
# Scans of the exposure inputs
# ---------------------------------------------------------------------------

def scan_counts(exposure_dir: Path) -> pd.DataFrame:
    """Feature counts per country × system × object_type from the exposure inputs."""
    rows = []
    for folder, system in FOLDERS.items():
        for f in sorted((exposure_dir / folder).glob(f"{folder}_*.parquet")):
            iso2 = f.stem.split("_")[-1]
            names = pq.ParquetFile(f).schema.names
            use = [c for c in ["osm_id", "object_type"] if c in names]
            t = pd.read_parquet(f, columns=use)
            g = t.groupby("object_type", dropna=False)
            r = pd.DataFrame({"n": g.size(),
                              "n_unique_osm": g.osm_id.nunique() if "osm_id" in t else pd.NA}).reset_index()
            r["system"], r["country"] = system, iso2
            rows.append(r)
            print(f"  {system} {iso2}: {len(t)}")
    return pd.concat(rows, ignore_index=True)


def scan_geometry_types(exposure_dir: Path) -> pd.Series:
    """Dominant geometry per (system, object_type), from the last block of each extraction report."""
    rows = []
    for f in (exposure_dir / "_stats").glob("*/report_*.txt"):
        folder, iso2 = f.stem.replace("report_", "").rsplit("_", 1)
        if folder not in FOLDERS:
            continue
        block = f.read_text(encoding="utf8", errors="ignore").split("===")[-1]
        for line in block.splitlines():
            m = re.match(r"\s*(\S+)\s+(Point|MultiPoint|LineString|MultiLineString|Polygon|MultiPolygon)\s+(\d+)\s", line)
            if m:
                rows.append(dict(system=FOLDERS[folder], object_type=m.group(1),
                                 gclass=m.group(2).replace("Multi", ""), n=int(m.group(3))))
    g = pd.DataFrame(rows).groupby(["system", "object_type", "gclass"]).n.sum().unstack(fill_value=0)
    return g.idxmax(axis=1)


# ---------------------------------------------------------------------------
# Tables
# ---------------------------------------------------------------------------

def table_s1(cnt: pd.DataFrame) -> pd.DataFrame:
    road_class = {k: v for k, v in C.ROAD_CLASS_MAP.items() if k.islower()}
    rows = []
    for s in SYS:
        present = cnt[cnt.system == s].groupby("object_type").n.sum()
        code = C.DICT_CIS_VULNERABILITY_FLOOD[CODE_KEY.get(s, s)]
        code_vals = [k for k in code if k == k.lower() and " " not in k]
        for v in sorted(set(present.index) | set(code_vals)):
            if s == "ports":
                tag = PORT_TAGS.get(v, f"? ({v})")
            elif s == "telecom":
                tag = TEL_TAGS.get(v, f"man_made={v}")
            else:
                tag = f"{KEYS[s]}={v}"
            rows.append(dict(system=SYS_LABEL[s], asset_type=v, osm_tag=tag,
                             analysis_class=road_class.get(v, v).replace("_", " ") if s == "roads" else v,
                             n_features=int(present.get(v, 0)), in_data=v in present.index, in_code=v in code_vals))
    return pd.DataFrame(rows)


def table_s2(cnt: pd.DataFrame) -> pd.DataFrame:
    p = cnt.pivot_table(index="country", columns="system", values="n", aggfunc="sum", fill_value=0)
    p = p.reindex(columns=SYS, fill_value=0)
    p.index = [f"{ISO2_NAME.get(c, c)} ({c})" for c in p.index]
    p = p.sort_index()
    p["Total"] = p.sum(axis=1)
    p.loc["Total"] = p.sum()
    p.columns = [SYS_LABEL.get(c, c) for c in p.columns]
    return p


def _ranges(ids) -> str:
    """Compress ['F7.4', 'F7.5', 'F7.6'] -> 'F7.4–F7.6'."""
    def key(x):
        m = re.match(r"([A-Z])(\d+)\.(\d+)(-C)?", x)
        return (m.group(1), int(m.group(2)), int(m.group(3)), m.group(4) or "")
    ids = sorted(set(ids), key=key)
    out, run = [], []
    for x in ids:
        if run:
            a, b = key(run[-1]), key(x)
            if a[:2] == b[:2] and a[3] == b[3] and b[2] == a[2] + 1:
                run.append(x)
                continue
            out.append(run[0] if len(run) == 1 else f"{run[0]}–{run[-1]}")
        run = [x]
    if run:
        out.append(run[0] if len(run) == 1 else f"{run[0]}–{run[-1]}")
    return ", ".join(out)


def table_s3(cnt: pd.DataFrame, fragility_path: Path, vulnerability_path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Curves actually applied per asset type (+ legend).

    Flood and windstorm: hazard_river.prepare_flood_curves / hazard_windstorm apply every
    curve listed for a system to every object type, minus run_pipeline.Config exclusions.
    Earthquake: the per-object-type list in constants.py (curves present in the fragility file).
    """
    eq_ids = set(pd.read_excel(fragility_path, "E_Frag_PGA", header=[0, 1]).columns.get_level_values(0))
    rows = []
    for s in SYS:
        k = CODE_KEY.get(s, s)
        present = set(cnt[cnt.system == s].object_type)
        fl = C.DICT_CIS_VULNERABILITY_FLOOD[k]
        wd = C.DICT_CIS_VULNERABILITY_WIND.get(k, {})
        eq = C.DICT_CIS_VULNERABILITY_EARTHQUAKE[k]
        fl_all = {c for v in fl.values() for c in v}
        wd_all = {c for v in wd.values() for c in v}
        fx = Config.FLOOD_CURVE_EXCLUSIONS.get(s, {})
        wx = Config.WIND_CURVE_EXCLUSIONS.get(s, {})
        for ot in sorted(present):
            f_eff = sorted(fl_all - set(fx.get(ot, []))) if ot in fl else []
            w_eff = sorted(wd_all - set(wx.get(ot, []))) if ot in wd else []
            e_eff = sorted(set(eq.get(ot, [])) & eq_ids)
            rows.append(dict(system=SYS_LABEL[s], asset_type=ot,
                             windstorm=_ranges(w_eff) or "none (no curve: no windstorm damage)",
                             flood_river_coastal=_ranges(f_eff) or "none",
                             earthquake=_ranges(e_eff) or "none",
                             code_list_windstorm=_ranges(wd.get(ot, [])),
                             code_list_flood=_ranges(fl.get(ot, []))))
    s3 = pd.DataFrame(rows)

    desc = {}
    for sheet in ["F_Vuln_Depth", "W_Vuln_V10m_3sec", "E_Frag_PGA"]:
        h = pd.read_excel(vulnerability_path, sheet, header=None, nrows=3)
        for j in range(1, h.shape[1]):
            i = h.iloc[0, j]
            if isinstance(i, str) and i not in desc:
                desc[i] = (str(h.iloc[1, j]), str(h.iloc[2, j]) if pd.notna(h.iloc[2, j]) else "", sheet)
    used = set()
    for s in SYS:
        k = CODE_KEY.get(s, s)
        for dct in (C.DICT_CIS_VULNERABILITY_FLOOD, C.DICT_CIS_VULNERABILITY_WIND, C.DICT_CIS_VULNERABILITY_EARTHQUAKE):
            for v in dct.get(k, {}).values():
                used |= set(v)
    leg = []
    for i in sorted(used, key=lambda x: (x[0], [int(t) for t in re.findall(r"\d+", x)])):
        if i in desc:
            t, d, sh = desc[i]
            src = "Nirandjan et al. (2024) database, sheet " + sh
        elif i.endswith("-C"):
            t, d = "Building (education/healthcare), parametric fragility", ""
            src = ("data/EQ_fragility.xlsx only (lognormal median/beta per damage state); "
                   "not in the Nirandjan et al. (2024) database")
        else:
            t, d, src = "", "", "not found in D2 or EQ_fragility.xlsx"
        leg.append(dict(id=i, hazard={"F": "Flood", "W": "Windstorm", "E": "Earthquake"}[i[0]],
                        infrastructure_type=t, details=d, source=src, in_EQ_fragility_file=i in eq_ids))
    return s3, pd.DataFrame(leg)


def table_s4(cnt: pd.DataFrame, geo: pd.Series) -> pd.DataFrame:
    rows = []
    for s in SYS:
        md = C.INFRASTRUCTURE_DAMAGE_VALUES[CODE_KEY.get(s, s)]
        for ot in sorted(set(cnt[cnt.system == s].object_type)):
            lo, mid, hi = md[ot]
            g = geo.get((s, ot), "")
            rows.append(dict(system=SYS_LABEL[s], asset_type=ot, geometry=g, unit=GEOM_UNIT.get(g, "?"),
                             min=lo, median=mid, max=hi))
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# Word export (optional, needs python-docx)
# ---------------------------------------------------------------------------

def write_docx(path: Path, s1, s2, s3, leg, s4) -> bool:
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        print("  python-docx not installed: CSVs written, no .docx")
        return False
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(8)

    def add_table(title, df, note=None):
        doc.add_paragraph().add_run(title).bold = True
        t = doc.add_table(rows=1, cols=len(df.columns))
        t.style = "Table Grid"
        for j, c in enumerate(df.columns):
            t.rows[0].cells[j].text = str(c)
            for r in t.rows[0].cells[j].paragraphs[0].runs:
                r.bold = True
        for _, row in df.iterrows():
            cells = t.add_row().cells
            for j, v in enumerate(row):
                cells[j].text = f"{v:,}" if isinstance(v, int) and not isinstance(v, bool) else str(v)
        if note:
            doc.add_paragraph(note)
        doc.add_paragraph()

    n_tags = int(s1.in_data.sum())
    s1d = (s1[s1.in_data].groupby("system", sort=False)
           .agg(osm_tags=("osm_tag", "; ".join), n_types=("asset_type", "count"), n_features=("n_features", "sum"))
           .reset_index())
    s1d["n_features"] = s1d.n_features.astype(int)
    add_table(f"Table S1. OpenStreetMap tags used for each infrastructure system ({n_tags} asset types; "
              "roads are analysed in 5 classes).",
              s1d.rename(columns={"system": "Infrastructure system", "osm_tags": "OSM tags (key=value)",
                                  "n_types": "Asset types", "n_features": "Features"}),
              "Roads: motorway, motorway_link, trunk and trunk_link are analysed as 'motorways and trunks'; "
              "primary(_link) as primary; secondary(_link) as secondary; tertiary(_link) as tertiary; residential, "
              "road, unclassified and track as other roads.")
    p2 = s2.reset_index().rename(columns={"index": "Country"})
    for c in p2.columns[1:]:
        p2[c] = p2[c].astype(int)
    add_table("Table S2. Number of critical infrastructure assets (OSM features) per country and system.", p2,
              "Features are OSM objects split at LAU boundaries. Zero cells are true zeros: Andorra, Liechtenstein "
              "and North Macedonia are landlocked (no ports); Andorra and Malta have no railway.")
    add_table("Table S3. Vulnerability (flood, windstorm) and fragility (earthquake) curve IDs applied per asset type.",
              s3[["system", "asset_type", "windstorm", "flood_river_coastal", "earthquake"]].rename(columns={
                  "system": "Infrastructure system", "asset_type": "Asset type", "windstorm": "Windstorm IDs",
                  "flood_river_coastal": "Flood IDs (river, coastal)", "earthquake": "Earthquake IDs"}),
              "IDs refer to the physical vulnerability database of Nirandjan et al. (2024): F = flood depth-damage "
              "curve, W = windstorm (3 s gust) damage curve, E = earthquake fragility curve (PGA) converted to a "
              f"damage ratio with damage-state ratios {HEQ.DAMAGE_RATIOS}. IDs with suffix '-C' are parametric "
              "building fragility curves (lognormal median and dispersion per damage state) that are not part of the "
              "Nirandjan et al. (2024) database; they are taken from [SOURCE TO BE CONFIRMED]. For flood and "
              "windstorm, every curve listed for a system is applied to every asset type in that system unless "
              "explicitly excluded; the table lists the curves actually applied.")
    add_table("Table S3 (legend). Description of curve IDs.",
              leg[["id", "hazard", "infrastructure_type", "details", "source"]])
    add_table("Table S4. Maximum reconstruction costs per asset type (min, median, max), from Nirandjan et al. (2024).",
              s4.rename(columns={"system": "Infrastructure system", "asset_type": "Asset type", "geometry": "Geometry",
                                 "unit": "Unit", "min": "Min", "median": "Median", "max": "Max"}),
              "Units follow the asset geometry: €/m for lines (damage = cost × exposed length), €/m² for polygons "
              "(× area) and €/unit for points. For railway windstorm damage the power-line cost (€/m) is used as a "
              "catenary proxy. Reference: Nirandjan S, Koks E E, Ye M, Pant R, van Ginkel K C H, Aerts J C J H and "
              "Ward P J 2024 Physical vulnerability database for critical infrastructure hazard risk assessments: a "
              "systematic review and data collection. Nat. Hazards Earth Syst. Sci. 24 4341–68. "
              "https://doi.org/10.5194/nhess-24-4341-2024")
    doc.save(path)
    return True


def main():
    parser = argparse.ArgumentParser(description="Build Supplementary Tables S1-S4")
    parser.add_argument("--exposure-dir", default=None, help="Exposure_files directory (default: config.yml)")
    parser.add_argument("--output-dir", default=None, help="Output directory (default: repo_root/tables)")
    args = parser.parse_args()

    exposure_dir = Path(args.exposure_dir) if args.exposure_dir else Path(C.load_config()["exposure_dir"])
    out = Path(args.output_dir) if args.output_dir else REPO_ROOT / "tables"
    out.mkdir(parents=True, exist_ok=True)

    print(f"Exposure inputs: {exposure_dir}")
    cnt = scan_counts(exposure_dir)
    geo = scan_geometry_types(exposure_dir)

    s1 = table_s1(cnt)
    s2 = table_s2(cnt)
    s3, leg = table_s3(cnt, REPO_ROOT / "data/EQ_fragility.xlsx",
                       REPO_ROOT / "data/Table_D2_Hazard_Fragility_and_Vulnerability_Curves_V1.1.0_conversions.xlsx")
    s4 = table_s4(cnt, geo)

    s1.to_csv(out / "S1_osm_tags.csv", index=False)
    s2.to_csv(out / "S2_asset_counts.csv")
    s3.to_csv(out / "S3_vulnerability_ids.csv", index=False)
    leg.to_csv(out / "S3_legend_ids.csv", index=False)
    s4.to_csv(out / "S4_max_damage.csv", index=False)
    write_docx(out / "Supplementary_Tables_S1-S4.docx", s1, s2, s3, leg, s4)

    print(f"S1: {int(s1.in_data.sum())} asset types in data | S2: {len(s2) - 1} countries, "
          f"{int(s2.loc['Total', 'Total']):,} features | tables written to {out}")


if __name__ == "__main__":
    main()
